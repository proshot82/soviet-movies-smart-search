# -*- coding: utf-8 -*-
"""Экспорт финального каталога в Excel (аналог полной базы для чтения/фильтрации).
Запуск из корня папки: python scripts/export_xlsx.py  → Госфильмофонд_каталог_v2.xlsx (Windows: python; контейнер: python3).
Листы: «Каталог» (все записи, 7 полей базы + сведения указателя + пометка реестра), «Реестр стабов», «Жанры», «Студии», «Сводка», «О файле».
Данные: state/catalog_final_v2.json, state/alignment_final.json, state/index_master_final.json, state/Ф5_стабы_реестр.json. Итоги на «Сводке»/«Жанрах»/«Студиях» — формулами."""
import json, hashlib, datetime, sys
from collections import Counter
from openpyxl import Workbook
from openpyxl.styles import Font, Alignment, PatternFill, Border, Side
from openpyxl.utils import get_column_letter
OUT = sys.argv[1] if len(sys.argv) > 1 else 'Госфильмофонд_каталог_v2.xlsx'
cat = json.load(open('state/catalog_final_v2.json', encoding='utf-8'))
al = json.load(open('state/alignment_final.json', encoding='utf-8'))
m = json.load(open('state/index_master_final.json', encoding='utf-8'))
reg = json.load(open('state/Ф5_стабы_реестр.json', encoding='utf-8'))
sha = hashlib.sha256(open('state/catalog_final_v2.json', 'rb').read()).hexdigest()
N = len(cat); assert N == len(al) and all(a['cat_pos'] == i and a['id'] == i + 1 for i, a in enumerate(al))
cards = {r['cat_pos']: r for r in m['rows'] if r['kind'] == 'card'}; assert len(cards) == N
regby = {}
for e in reg['entries']: regby.setdefault(e['cat'], []).append(e)
VOLS = [f'v{i}_' + s for i, s in zip(range(1, 16), ['1918-1935', '1936-1957', '1958-1963', '1964-1965', '1966-1967', '1968-1969', '1970-1971', '1972-1973', '1974-1975', '1976-1977', '1978-1979', '1980-1981', '1982-1983', '1984-1985', '1986-1987'])]
assert set(r['source'] for r in cat) == set(VOLS)
F = 'Arial'
fN = Font(name=F, size=10); fB = Font(name=F, size=10, bold=True); fH = Font(name=F, size=10, bold=True, color='FFFFFF'); fT = Font(name=F, size=12, bold=True); fI = Font(name=F, size=9, italic=True, color='555555')
fillH = PatternFill('solid', fgColor='305496'); fillS = PatternFill('solid', fgColor='DDEBF7')
thin = Side(style='thin', color='BFBFBF'); box = Border(left=thin, right=thin, top=thin, bottom=thin)
top = Alignment(vertical='top'); wrap = Alignment(vertical='top', wrap_text=True); ctr = Alignment(horizontal='center', vertical='top')
wb = Workbook()
# ── Каталог ──
ws = wb.active; ws.title = 'Каталог'
H = ['id', 'Название (title)', 'Жанр (genre)', 'Год (year)', 'Студия (studio)', 'Аннотация (annotation)', 'Том (source)', '№ по указателю', 'Название по указателю', 'Другие названия (указатель)', 'Строка xlsx-указателя', 'Реестр стабов (пустые/неполные поля)']
W = [7, 42, 22, 7, 34, 100, 14, 10, 40, 30, 10, 34]
ws.append(H)
for c in range(1, len(H) + 1):
    cell = ws.cell(row=1, column=c); cell.font = fH; cell.fill = fillH; cell.alignment = Alignment(vertical='center', wrap_text=True, horizontal='center'); cell.border = box
    ws.column_dimensions[get_column_letter(c)].width = W[c - 1]
ws.row_dimensions[1].height = 30
for i, r in enumerate(cat):
    a = al[i]; cd = cards[i]
    alt = '; '.join(cd.get('alt_titles') or [])
    rg = '; '.join(f"{e['field']}: {e['state']}" for e in regby.get(i, []))
    row = [r['id'], r['title'], r['genre'], r['year'], r['studio'], r['annotation'], r['source'], a['idx_num'], cd['title_canon'], alt or None, a['xlsx_row'], rg or None]
    ws.append(row)
    rr = i + 2
    for c in range(1, len(H) + 1):
        cell = ws.cell(row=rr, column=c); cell.font = fN
        cell.alignment = wrap if c in (2, 5, 6, 9, 10, 12) else (ctr if c in (1, 4, 8, 11) else top)
    ws.cell(row=rr, column=1).number_format = '0'; ws.cell(row=rr, column=4).number_format = '0'; ws.cell(row=rr, column=11).number_format = '0'
    if isinstance(a['idx_num'], int): ws.cell(row=rr, column=8).number_format = '0'
ws.freeze_panes = 'C2'; ws.auto_filter.ref = f'A1:{get_column_letter(len(H))}{N + 1}'
LAST = N + 1
# ── Реестр стабов ──
wr = wb.create_sheet('Реестр стабов')
HR = ['cat (индекс 0..N-1)', 'id', 'Название', 'Том', 'Поле', 'Состояние', 'Основание', 'Происхождение', 'Прежний cat (Ф5)']
WR = [10, 40, 14, 12, 14, 90, 22, 10]
wr.append(HR)
for c in range(1, len(HR) + 1):
    cell = wr.cell(row=1, column=c); cell.font = fH; cell.fill = fillH; cell.alignment = Alignment(vertical='center', wrap_text=True, horizontal='center'); cell.border = box
    wr.column_dimensions[get_column_letter(c)].width = ([8] + WR)[c - 1]
for e in sorted(reg['entries'], key=lambda e: (e['cat'], e['field'])):
    wr.append([e['cat'], e['id'], e['title'], e['source'], e['field'], e['state'], e['reason'], e.get('origin'), e.get('cat_f5')])
for row in wr.iter_rows(min_row=2, max_row=wr.max_row):
    for cell in row: cell.font = fN; cell.alignment = wrap if cell.column in (3, 7) else top
wr.freeze_panes = 'C2'; wr.auto_filter.ref = f'A1:{get_column_letter(len(HR))}{wr.max_row}'
# ── Жанры / Студии (формулы EXACT — с учётом регистра) ──
def list_sheet(name, col, values):
    w = wb.create_sheet(name)
    w.append([name[:-1] if name.endswith('ы') else name, 'Фильмов']); w.append(['(всего уникальных:', f'=COUNTA(A3:A{len(values) + 2})'])
    for c in (1, 2):
        cell = w.cell(row=1, column=c); cell.font = fH; cell.fill = fillH; cell.border = box; cell.alignment = Alignment(horizontal='center')
    w.cell(row=2, column=1).font = fI; w.cell(row=2, column=2).font = fB
    for k, v in enumerate(values, start=3):
        w.cell(row=k, column=1, value=v).font = fN
        f = w.cell(row=k, column=2, value=f"=SUMPRODUCT(--EXACT('Каталог'!${col}$2:${col}${LAST},A{k}))"); f.font = fN; f.number_format = '0'
    w.column_dimensions['A'].width = 60; w.column_dimensions['B'].width = 10; w.freeze_panes = 'A3'
    return w
gc = Counter(r['genre'] for r in cat if r['genre']); sc = Counter(r['studio'] for r in cat if r['studio'])
list_sheet('Жанры', 'C', [g for g, _ in gc.most_common()])
list_sheet('Студии', 'E', [s for s, _ in sc.most_common()])
# ── Сводка ──
wsv = wb.create_sheet('Сводка')
wsv.column_dimensions['A'].width = 46; wsv.column_dimensions['B'].width = 16; wsv.column_dimensions['C'].width = 70
def put(r, c, v, font=fN, fmt=None, align=None):
    cell = wsv.cell(row=r, column=c, value=v); cell.font = font
    if fmt: cell.number_format = fmt
    if align: cell.alignment = align
    return cell
put(1, 1, 'Каталог Госфильмофонда — сводка', fT)
put(2, 1, 'Аннотированный каталог советских художественных фильмов 1918–1987 (15 томов); финальная версия v2 (Ф6-П3в), собрана ' + datetime.date.today().strftime('%d.%m.%Y') + '.', fI)
put(3, 1, 'Источник: state/catalog_final_v2.json, sha256 ' + sha, fI)
put(5, 1, 'Записей всего', fB); put(5, 2, f"=COUNTA('Каталог'!A2:A{LAST})", fB, '0')
put(6, 1, 'Уникальных жанров', fN); put(6, 2, "=Жанры!B2", fN, '0')
put(7, 1, 'Уникальных студий', fN); put(7, 2, "=Студии!B2", fN, '0')
put(8, 1, 'Пустой жанр (None)', fN); put(8, 2, f"=COUNTBLANK('Каталог'!C2:C{LAST})", fN, '0'); put(8, 3, 'причины — лист «Реестр стабов»', fI)
put(9, 1, 'Пустая студия (None)', fN); put(9, 2, f"=COUNTBLANK('Каталог'!E2:E{LAST})", fN, '0')
put(10, 1, 'Пустая аннотация (None)', fN); put(10, 2, f"=COUNTBLANK('Каталог'!F2:F{LAST})", fN, '0')
put(11, 1, 'Аннотаций короче 100 знаков (непустых)', fN); put(11, 2, f"=SUMPRODUCT((LEN('Каталог'!F2:F{LAST})>0)*(LEN('Каталог'!F2:F{LAST})<100))", fN, '0'); put(11, 3, 'в основном однофразные логлайны и формулы источника (в печати прозы нет)', fI)
put(12, 1, 'Средняя длина аннотации, знаков', fN); put(12, 2, f"=ROUND(SUMPRODUCT(LEN('Каталог'!F2:F{LAST}))/COUNTA('Каталог'!F2:F{LAST}),0)", fN, '0')
r0 = 14
put(r0, 1, 'По томам (source)', fB); put(r0, 2, 'Записей', fB); put(r0, 3, 'Годы производства (по тому)', fB)
for k, v in enumerate(VOLS):
    put(r0 + 1 + k, 1, v); put(r0 + 1 + k, 2, f"=COUNTIF('Каталог'!$G$2:$G${LAST},A{r0 + 1 + k})", fN, '0'); put(r0 + 1 + k, 3, v.split('_')[1].replace('-', '–'), fI)
put(r0 + 16, 1, 'Итого', fB); put(r0 + 16, 2, f"=SUM(B{r0 + 1}:B{r0 + 15})", fB, '0')
r1 = r0 + 18
put(r1, 1, 'По десятилетиям (год производства)', fB); put(r1, 2, 'Записей', fB)
decs = list(range(1910, 1990, 10))
for k, d in enumerate(decs):
    put(r1 + 1 + k, 1, f'{d}–{d + 9}'); put(r1 + 1 + k, 2, f"=COUNTIFS('Каталог'!$D$2:$D${LAST},\">=\"&{d},'Каталог'!$D$2:$D${LAST},\"<=\"&{d + 9})", fN, '0')
put(r1 + 1 + len(decs), 1, 'Итого', fB); put(r1 + 1 + len(decs), 2, f"=SUM(B{r1 + 1}:B{r1 + len(decs)})", fB, '0')
put(r1 + 2 + len(decs), 1, 'Год min / max', fN); put(r1 + 2 + len(decs), 2, f"=MIN('Каталог'!D2:D{LAST})&\" / \"&MAX('Каталог'!D2:D{LAST})", fN)
r2 = r1 + 4 + len(decs)
dups = [(k, [r['id'] for r in cat if (r['title'], r['year']) == k]) for k, c in Counter((r['title'], r['year']) for r in cat).items() if c > 1]
put(r2, 1, 'Точные совпадения «название + год» (разные фильмы разных студий)', fB); put(r2, 2, 'Записей', fB); put(r2, 3, 'id и студии', fB)
for k, ((t, y), ids) in enumerate(sorted(dups, key=lambda x: x[1][0])):
    put(r2 + 1 + k, 1, f'«{t}», {y}'); put(r2 + 1 + k, 2, f"=COUNTIFS('Каталог'!$B$2:$B${LAST},\"{t}\",'Каталог'!$D$2:$D${LAST},{y})", fN, '0')
    put(r2 + 1 + k, 3, '; '.join(f"id {i}: {cat[i - 1]['studio']}" for i in ids), fI)
put(r2 + 2 + len(dups), 1, 'Совпадений «название + год + студия» нет (проверено при сборке).', fI)
for row in wsv.iter_rows(min_row=1, max_row=wsv.max_row):
    for cell in row:
        if cell.column == 3: cell.alignment = wrap
# ── О файле ──
wo = wb.create_sheet('О файле'); wo.column_dimensions['A'].width = 34; wo.column_dimensions['B'].width = 120
rows = [
    ('Что это', 'Полный аннотированный каталог Госфильмофонда (советские художественные фильмы 1918–1987, 15 печатных томов) в виде таблицы — тот же набор записей и полей, что и в базе поискового движка (таблица films).'),
    ('Версия', f'catalog_final_v2.json, {N} записей, sha256 {sha}; экспорт {datetime.date.today().strftime("%d.%m.%Y")} скриптом scripts/export_xlsx.py.'),
    ('id', 'Сквозной номер записи 1..N в порядке томов и статей (совпадает с id в базе движка при импорте v2).'),
    ('Название (title)', 'Каноническое написание названия (сверено с алфавитным указателем тома); варианты и подзаголовки — в скобках, как в печати.'),
    ('Жанр (genre)', 'Жанр из печатной головы статьи в полной форме (сокращения тома раскрыты); пусто — жанр в печати отсутствует (см. «Реестр стабов»).'),
    ('Год (year)', 'Год производства (не год выпуска на экран).'),
    ('Студия (studio)', 'Студия(и) по печатной голове статьи, без метража и числа частей; пусто — в печати отсутствует.'),
    ('Аннотация (annotation)', 'Проза аннотации из печатного тома целиком: логлайн, сюжет, примечания (источник, «Фильм не сохранился», «В РСФСР не демонстрировался» и т. п.). Титры (авторы, роли, исполнители, благодарности), библиография и наградные блоки в аннотацию не входят — так задан канон базы. Пусто — в печати аннотации нет.'),
    ('Том (source)', 'Метка тома: vN_ГГГГ-ГГГГ (номер тома и охватываемые годы производства).'),
    ('№ по указателю', 'Номер статьи по алфавитному указателю тома (печатная нумерация; литерные номера вида «1628а» — текстом).'),
    ('Название по указателю', 'Написание названия по мастер-указателю (сводный именной референс «фильм-год», построенный из указателей всех томов).'),
    ('Другие названия (указатель)', 'Альтернативные названия из указателя (через «; »), если есть.'),
    ('Строка xlsx-указателя', 'Номер строки в файле «Госфильмофонд именной указатель 19181987.xlsx» (для сверки).'),
    ('Реестр стабов', 'Пометка для записей с пустыми или заведомо неполными полями: «поле: состояние» (none — в печати нет; source-only — в статье только формула источника; formula — аннотация-формула; content/note/partial — см. основание на листе «Реестр стабов»).'),
    ('Листы', '«Каталог» — все записи (фильтры по заголовкам, закреплены id и название); «Реестр стабов» — основания для пустых/неполных полей; «Жанры», «Студии» — частоты (формулы); «Сводка» — итоги формулами по листу «Каталог».'),
    ('Кодировка/особенности', 'Буква «ё» в томах не печатается (везде «е», кроме единичных случаев печати). Четыре аннотации многоновелльных фильмов содержат переводы строк. Написание сохраняет орфографию тома (например, «Попраный обет»).'),
]
wo.append(['Поле / раздел', 'Описание'])
for c in (1, 2):
    cell = wo.cell(row=1, column=c); cell.font = fH; cell.fill = fillH; cell.border = box
for k, v in rows:
    wo.append([k, v])
for row in wo.iter_rows(min_row=2, max_row=wo.max_row):
    row[0].font = fB; row[0].alignment = top; row[1].font = fN; row[1].alignment = wrap
wb.save(OUT)
print('saved', OUT, 'records', N, 'genres', len(gc), 'studios', len(sc), 'registry', len(reg['entries']))
