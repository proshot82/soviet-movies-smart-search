# -*- coding: utf-8 -*-
"""Пересчёт формул xlsx через LibreOffice (headless): calculateAll + store, затем проверка на ошибки формул.
Запуск: python3 scripts/xlsx_recalc.py файл.xlsx [timeout]"""
import os, sys, subprocess, tempfile, time, json
from pathlib import Path
from openpyxl import load_workbook
MACRO = """<?xml version="1.0" encoding="UTF-8"?>
<!DOCTYPE script:module PUBLIC "-//OpenOffice.org//DTD OfficeDocument 1.0//EN" "module.dtd">
<script:module xmlns:script="http://openoffice.org/2000/script" script:name="Module1" script:language="StarBasic">
    Sub RecalculateAndSave()
      ThisComponent.calculateAll()
      ThisComponent.store()
      ThisComponent.close(True)
    End Sub
</script:module>"""
fn = sys.argv[1]; timeout = int(sys.argv[2]) if len(sys.argv) > 2 else 150
absp = str(Path(fn).absolute()); env = dict(os.environ, SAL_USE_VCLPLUGIN='svp')
with tempfile.TemporaryDirectory(prefix='lo-prof-') as prof:
    url = Path(prof).as_uri()
    subprocess.run(['soffice', '--headless', '--terminate_after_init', f'-env:UserInstallation={url}'], capture_output=True, timeout=timeout, env=env)
    md = Path(prof) / 'user' / 'basic' / 'Standard'; assert md.exists(), 'no LO profile'
    (md / 'Module1.xba').write_text(MACRO)
    before = (os.stat(absp).st_mtime_ns, os.stat(absp).st_size); t0 = time.time()
    r = subprocess.run(['soffice', '--headless', '--norestore', f'-env:UserInstallation={url}', 'vnd.sun.star.script:Standard.Module1.RecalculateAndSave?language=Basic&location=application', absp], capture_output=True, text=True, timeout=timeout, env=env)
    assert r.returncode == 0, r.stderr
    assert (os.stat(absp).st_mtime_ns, os.stat(absp).st_size) != before, 'file not rewritten'
wb = load_workbook(fn, data_only=True); errs = {}; nform = 0
wbf = load_workbook(fn)
for name in wbf.sheetnames:
    for row in wbf[name].iter_rows():
        for c in row:
            if isinstance(c.value, str) and c.value.startswith('='): nform += 1
for name in wb.sheetnames:
    for row in wb[name].iter_rows():
        for c in row:
            if isinstance(c.value, str) and any(e in c.value for e in ('#VALUE!', '#DIV/0!', '#REF!', '#NAME?', '#NULL!', '#NUM!', '#N/A')): errs.setdefault(c.value[:8], []).append(f'{name}!{c.coordinate}')
print(json.dumps({'status': 'errors_found' if errs else 'success', 'total_formulas': nform, 'total_errors': sum(len(v) for v in errs.values()), 'errors': {k: v[:20] for k, v in errs.items()}, 'seconds': round(time.time() - t0, 1)}, ensure_ascii=False))
