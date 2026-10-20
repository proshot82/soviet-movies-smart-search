#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""Список работ по сюжетным аннотациям (этап 2 плана) и применение заполненного списка.

    python tools/annotation_worklist.py export --estimate                 # сколько стоит классификация
    python tools/annotation_worklist.py export --classify --out imports/аннотации_список_работ.xlsx
    python tools/annotation_worklist.py apply imports/аннотации_список_работ.xlsx --out imports/аннотации_правки.json
    python film_search.py import imports/аннотации_правки.json --dry-run
    python film_search.py import imports/аннотации_правки.json
    python film_search.py embed --resume

export — xlsx по фильмам film_id ≥ 7175 (аннотации каталогов ГФФ 1–7174 закрыты). Порядок:
  1) аннотация обрывается на «…» / «...»;
  2) «не сюжет» — текст не говорит, что происходит в фильме (только о литературной основе,
     судьбе фильма, актёрах, жанре). Определяет дешёвая модель (FLAG_MODEL), ответы кешируются
     в imports/аннотации_классы.json по хешу аннотации — повторный запуск платит только за новые;
  3) остальные — от коротких к длинным.
  Колонки «Новая аннотация» и «Источник новой аннотации» заполняются вручную. С --gff они заранее
  заполняются описанием с gosfilmofond.ru (tools/gff_site_catalog.py), если оно лучше текущего; как найдено
  совпадение — в колонке «Совпадение» (год ±1 — проверить). С --wiki — так же разделом «Сюжет» из Википедии
  (tools/wiki_plots.py) там, где нет предложения ГФФ. Каждую пару сверяет модель («тот же фильм?»), решения
  кешируются в imports/аннотации_сверка.json. Ненужное предложение просто стереть.

apply — заполненный xlsx → JSON-правки [{"film_id", "annotation", "source"}] для film_search.py import.
  Ошибка в любой строке (film_id < 7175 или нет в базе, название не совпадает с базой — строки
  сдвинуты, нет источника, обрыв на многоточие, film_id дважды) — JSON не пишется.
"""
import argparse
import gzip
import hashlib
import json
import re
import sqlite3
import sys
import urllib.parse
from pathlib import Path

import openpyxl
from openpyxl.styles import Alignment, Font

ROOT = Path(__file__).resolve().parent.parent
sys.path.insert(0, str(ROOT))
sys.path.insert(0, str(ROOT / "tools"))
from film_search import DB_PATH, get_openai_client  # noqa: E402

FIRST_OPEN_ID = 7175        # film_id 1–7174 — каталоги ГФФ, аннотации не трогаем (PLAN.md, раздел 2)
SHORT_CHARS = 300
CLASSES_PATH = ROOT / "imports" / "аннотации_классы.json"
MATCHES_PATH = ROOT / "imports" / "аннотации_сверка.json"     # «тот же фильм?» по парам (film_id, текст источника)

# Проба на 50 + 8 трудных аннотациях 19.10.2026: gpt-5.4-nano и с "none", и с "low" считал сюжетом
# «Семнадцать мгновений весны» (8327: только время действия и слава фильма), gpt-5.4-mini — нет
FLAG_MODEL = "gpt-5.4-mini"
FLAG_EFFORT = "low"
FLAG_BATCH = 25             # аннотаций в одном запросе
FLAG_CLIP = 1500            # сколько текста видит сверка «тот же фильм?»
# Оценка стоимости: цены gpt-5.4-mini $0,75 / $4,50 за 1 млн входных / выходных токенов — по сводкам
# цен на 19.10.2026, на странице OpenAI не проверены. ~2 символа на токен — по эмбеддингам этих
# же аннотаций (2,28 млн символов чанков = 1,16 млн токенов); выходных ~30 на аннотацию — на пробе
# (ответ + рассуждение при "low").
PRICE_IN, PRICE_OUT = 0.75, 4.50
CHARS_PER_TOKEN = 2.0
PROMPT_TOKENS, ANSWER_TOKENS = 450, 30      # на запрос / на одну аннотацию в ответе

FLAG_PROMPT = """Ниже аннотации советских фильмов, телефильмов, телеспектаклей и мультфильмов. Для каждой \
определи, насколько полно она пересказывает СЮЖЕТ:

"full" — сюжет рассказан до развязки: понятно, чем кончается история (как разрешился конфликт, что стало \
с героями). Для концерта, документального или научно-популярного фильма — конкретно перечислено, что показано.
"setup" — есть герои и завязка (хотя бы одно событие), но не сказано, чем всё кончается: анонс, обрыв на \
«…и тут начинается…», «но однажды…».
"none" — сюжета нет: только литературная основа («По одноимённой пьесе …»), жанр, общая тема («о любви к \
природе»), время и место действия, судьба и слава фильма, актёры, театр, награды.

Верни СТРОГО JSON-массив без пояснений: [{{"id": <номер>, "level": "full"|"setup"|"none"}}, ...] — по одному \
элементу на каждую аннотацию.

{items}"""

HEADERS = ["Приоритет", "Группа", "film_id", "Название", "Год", "Тип", "Студия", "Длина", "Обрыв",
           "Полнота", "Текущая аннотация", "Текущий источник", "Поиск в Википедии",
           "Новая аннотация", "Источник новой аннотации", "Совпадение"]
WIDTHS = [9, 14, 8, 36, 6, 15, 24, 7, 7, 15, 70, 30, 14, 70, 40, 30]
# Страны каталога ГФФ, при которых фильм может быть советским (остальные — зарубежные, не сопоставляются)
SOVIET_COUNTRY = re.compile(r"ссср|cccр|россия|рсфср|украин|белорус|беларус|грузи|армени|азербайджан|казах|узбек|"
                            r"киргиз|кыргыз|таджик|туркмен|молдав|молдов|латви|литв|эстони", re.I)


def ann_hash(text):
    return hashlib.sha1((text or "").encode("utf-8")).hexdigest()[:16]


def is_truncated(text):
    return bool(re.search(r"(…|\.\.\.)\s*$", text or ""))


def load_films(db):
    conn = sqlite3.connect(f"file:{db}?mode=ro", uri=True)
    try:
        rows = conn.execute("""SELECT film_id, title, year, kind, studio, annotation, source FROM films
                               WHERE film_id >= ? ORDER BY film_id""", (FIRST_OPEN_ID,)).fetchall()
    finally:
        conn.close()
    keys = ("film_id", "title", "year", "kind", "studio", "annotation", "source")
    return [dict(zip(keys, r)) for r in rows]


def load_classes():
    if CLASSES_PATH.exists():
        return {int(k): v for k, v in json.loads(CLASSES_PATH.read_text(encoding="utf-8")).items()}
    return {}


def save_classes(classes):
    CLASSES_PATH.parent.mkdir(parents=True, exist_ok=True)
    data = {str(k): classes[k] for k in sorted(classes)}
    CLASSES_PATH.write_text(json.dumps(data, ensure_ascii=False, indent=0), encoding="utf-8")


LEVELS = ("full", "setup", "none")      # полный сюжет / только завязка / не сюжет
LEVEL_NAMES = {"full": "полный сюжет", "setup": "только завязка", "none": "не сюжет"}


def cached_level(classes, film):
    """"full" / "setup" / "none" из кеша; None — не классифицирован, аннотация изменилась или запись старого
    формата (только «сюжет / не сюжет», до 20.10.2026)."""
    c = classes.get(film["film_id"])
    return c.get("level") if c and c.get("hash") == ann_hash(film["annotation"]) else None


def cached_plot(classes, film):
    """True / False из кеша, None — если фильм не классифицирован или аннотация с тех пор изменилась."""
    c = classes.get(film["film_id"])
    return c["plot"] if c and c.get("hash") == ann_hash(film["annotation"]) else None


def estimate(films):
    chars = sum(len(head_tail(f["annotation"])) for f in films)
    batches = -(-len(films) // FLAG_BATCH)
    t_in = chars / CHARS_PER_TOKEN + batches * PROMPT_TOKENS
    t_out = len(films) * ANSWER_TOKENS
    return t_in, t_out, (t_in * PRICE_IN + t_out * PRICE_OUT) / 1e6


def head_tail(text, head=900, tail=600):
    """Начало и конец длинного текста: развязка — в конце, модели её надо видеть."""
    text = text or ""
    return text if len(text) <= head + tail else text[:head] + " […] " + text[-tail:]


def classify(films, classes):
    """Классифицировать аннотации моделью FLAG_MODEL; кеш сохраняется после каждого запроса."""
    client = get_openai_client()
    used_in = used_out = 0
    for start in range(0, len(films), FLAG_BATCH):
        batch = films[start:start + FLAG_BATCH]
        items = "\n\n".join(f"[{f['film_id']}] «{f['title']}» ({f['year']}): {head_tail(f['annotation'])}"
                            for f in batch)
        answer = None
        for attempt in range(3):
            try:
                resp = client.chat.completions.create(
                    model=FLAG_MODEL,
                    messages=[{"role": "user", "content": FLAG_PROMPT.format(items=items)}],
                    reasoning_effort=FLAG_EFFORT,
                    max_completion_tokens=8000,
                )
                used_in += resp.usage.prompt_tokens
                used_out += resp.usage.completion_tokens
                text = resp.choices[0].message.content.strip().replace("```json", "").replace("```", "")
                answer = {int(it["id"]): it["level"] for it in json.loads(text) if it.get("level") in LEVELS}
                break
            except Exception as e:      # сбой сети или формата ответа — повтор
                print(f"  запрос {start // FLAG_BATCH + 1}: {e}", file=sys.stderr)
        if answer is None:
            print(f"  film_id {batch[0]['film_id']}–{batch[-1]['film_id']} не классифицированы", file=sys.stderr)
            continue
        for f in batch:
            if f["film_id"] in answer:
                classes[f["film_id"]] = {"level": answer[f["film_id"]], "plot": answer[f["film_id"]] != "none",
                                         "hash": ann_hash(f["annotation"]), "model": FLAG_MODEL}
        save_classes(classes)
        done = min(start + FLAG_BATCH, len(films))
        if done % 500 < FLAG_BATCH or done == len(films):
            print(f"  классифицировано {done}/{len(films)}")
    cost = (used_in * PRICE_IN + used_out * PRICE_OUT) / 1e6
    print(f"Токенов: {used_in:,} входных, {used_out:,} выходных; ≈ ${cost:.3f}")


def better(film, new, classes):
    """Новый текст лучше текущей аннотации: не обрыв, длиннее старой, не её начало; если старая не обрыв и не
    «не сюжет» — длиннее хотя бы на 40 символов."""
    from xlsx_to_json import norm
    old = film["annotation"] or ""
    if not new or is_truncated(new) or len(new) <= len(old) or norm(old).startswith(norm(new)):
        return False
    weak = is_truncated(old) or cached_plot(classes, film) is False or norm(new).startswith(norm(old).rstrip(". "))
    return weak or len(new) >= len(old) + 40


SERIES_PATH = ROOT / "imports" / "аннотации_сериалы.json"
SERIES_PROMPT = """Ниже разделы «Сюжет» статей Википедии о советских сериалах, разбитые по сериям. Для каждого \
определи, как устроен сериал: "through" — сквозной сюжет (серии продолжают одну историю: «Семнадцать мгновений \
весны», «Гостья из будущего»), "anthology" — серии — отдельные законченные истории с общими героями («Следствие \
ведут знатоки», «Приключения Шерлока Холмса» по отдельным рассказам).

Верни СТРОГО JSON-массив без пояснений: [{{"id": <номер>, "kind": "through"|"anthology"}}, ...]

{items}"""


def series_kinds(recs):
    """{film_id: "through"|"anthology"} для посерийных сюжетов Википедии; кеш — SERIES_PATH (там же ручные
    решения; "skip" — текст не брать, например внутри сюжета списки «В ролях»)."""
    cache = json.loads(SERIES_PATH.read_text(encoding="utf-8")) if SERIES_PATH.exists() else {}
    todo = [r for r in recs if str(r["film_id"]) not in cache]
    if todo:
        client = get_openai_client()
        for start in range(0, len(todo), 10):
            batch = todo[start:start + 10]
            items = "\n\n".join(f"[{r['film_id']}] «{r['title']}» ({r['year']}):\n{r['plot'][:2500]}" for r in batch)
            for attempt in range(3):
                try:
                    resp = client.chat.completions.create(
                        model=FLAG_MODEL, messages=[{"role": "user", "content": SERIES_PROMPT.format(items=items)}],
                        reasoning_effort=FLAG_EFFORT, max_completion_tokens=8000)
                    text = resp.choices[0].message.content.strip().replace("```json", "").replace("```", "")
                    for it in json.loads(text):
                        if it.get("kind") in ("through", "anthology"):
                            cache[str(int(it["id"]))] = it["kind"]
                    break
                except Exception as e:      # сбой сети или формата ответа — повтор
                    print(f"  сериалы: {e}", file=sys.stderr)
            SERIES_PATH.write_text(json.dumps(cache, ensure_ascii=False, indent=0), encoding="utf-8")
    return {r["film_id"]: cache.get(str(r["film_id"])) for r in recs}


def wiki_suggestions(films, classes, path, taken):
    """{film_id: (текст, url, примечание)} — раздел «Сюжет» Википедии (tools/wiki_plots.py), если он лучше
    текущей аннотации и для фильма нет предложения ГФФ (taken). Посерийный сюжет сериала со сквозным сюжетом
    не предлагается — нужен общий сюжет (правило владельца); посерийный сюжет сериала из отдельных историй — да."""
    by_id = {f["film_id"]: f for f in films}
    recs = []
    with gzip.open(path, "rt", encoding="utf-8") as f:
        for line in f:
            r = json.loads(line)
            # ГФФ в приоритете: текст, уже взятый с gosfilmofond.ru, Википедией не заменяется
            if (r.get("plot") and r["film_id"] in by_id and r["film_id"] not in taken
                    and "gosfilmofond.ru" not in (by_id[r["film_id"]]["source"] or "")):
                recs.append(r)
    kinds = series_kinds([r for r in recs if r["episodes"]])
    out, skipped_series = {}, 0
    for r in recs:
        film = by_id[r["film_id"]]
        if r["episodes"] and kinds.get(r["film_id"]) != "anthology":
            skipped_series += 1
            continue
        new = clean_annotation(r["plot"])
        if better(film, new, classes):
            out[r["film_id"]] = (new, r["url"], f"Википедия: «{r['page']}»"
                                 + (" — посерийно, сериал из отдельных историй" if r["episodes"] else ""))
    return out, skipped_series


# Пересказ из интернета не берётся, если в нём упоминаются источники или остался мусор страниц
WEB_BAD = re.compile(r"подробнее\s*>>|в текущей аннотации|в карточке|по источникам|по данным источник|википеди|"
                     r"imdb|likefilm|кино-театр\.ру|kino-teatr|кинопоиск|источник[и]? не (сообщают|раскрывают)|"
                     r"не раскрывают|не сообщают|над[её]жн\w* (источник|описани|данн)", re.I)


def web_problem(film, rec):
    """Причина отказа для пересказа из интернета или None. Пакет помнит хеш аннотации, которую видел агент:
    если аннотация в базе с тех пор изменилась, пересказ не применяется."""
    text = rec.get("annotation") or ""
    if rec.get("ann_hash") and rec["ann_hash"] != hashlib.sha1((film["annotation"] or "").encode()).hexdigest()[:16]:
        return "аннотация в базе изменилась после выдачи пакета"
    if not [u for u in rec.get("sources") or [] if str(u).startswith("http")]:
        return "нет URL источника"
    if len(text) < 80:
        return "короче 80 знаков"
    letters = [ch for ch in text if ch.isalpha()]
    if not letters or sum("а" <= ch.lower() <= "я" or ch in "ёЁ" for ch in letters) / len(letters) < 0.5:
        return "меньше половины кириллицы"
    m = WEB_BAD.search(text)
    if m:
        return f"в тексте «{m.group(0)}»"
    return None


def web_suggestions(films, classes, paths, taken, batches=(), skip_path=None):
    """{film_id: (текст, источники, примечание)} — пересказы, найденные агентами в интернете (JSON Lines:
    film_id, status full|partial|keep|not_found, annotation, sources, identity, note). Тождество фильма проверял
    агент (identity), модель-сверка их не перепроверяет. ГФФ и Википедия в приоритете (taken).
    batches — входные пакеты (JSON): из них берётся хеш аннотации, которую видел агент."""
    by_id = {f["film_id"]: f for f in films}
    seen_hash = {}
    for b in batches:
        for f in json.loads(Path(b).read_text(encoding="utf-8")):
            seen_hash[f["film_id"]] = f.get("ann_hash")
    skip = json.loads(Path(skip_path).read_text(encoding="utf-8")) if skip_path else {}   # {film_id: причина}
    out, rejected = {}, []
    for path in paths:
        opener = gzip.open if str(path).endswith(".gz") else open
        with opener(path, "rt", encoding="utf-8") as f:
            for line in f:
                if not line.strip():
                    continue
                r = json.loads(line)
                if "status" not in r:          # старый формат (found / plot)
                    r = dict(r, status=("full" if r.get("complete") else "partial") if r.get("found") else "not_found",
                             annotation=r.get("plot"))
                film = by_id.get(r.get("film_id"))
                if not film or r.get("status") not in ("full", "partial") or r["film_id"] in taken:
                    continue
                if str(r["film_id"]) in skip:
                    rejected.append((r["film_id"], "исключён при проверке отчёта: " + skip[str(r["film_id"])]))
                    continue
                r.setdefault("ann_hash", seen_hash.get(r["film_id"]))
                why = web_problem(film, r)
                if why:
                    rejected.append((r["film_id"], why))
                    continue
                new = clean_annotation(r["annotation"])
                if better(film, new, classes) or r["status"] == "full" and cached_level(classes, film) != "full":
                    srcs = [u for u in r.get("sources") or [] if str(u).startswith("http")]
                    out[r["film_id"]] = (new, "; ".join(srcs), "поиск в интернете"
                                         + ("" if r["status"] == "full" else " — без развязки"))
    if rejected:
        print(f"  пересказов из интернета отклонено проверками: {len(rejected)}")
        for fid, why in rejected[:20]:
            print(f"    {fid}: {why}")
    return out


def gff_suggestions(films, classes, path):
    """{film_id: (текст, url, примечание)} — описания каталога gosfilmofond.ru, которые лучше текущей аннотации.
    Совпадение: то же название (целиком или без уточнения в скобках) и год ±1; при нескольких кандидатах
    берётся единственный с точным годом, иначе фильм пропускается. Телеспектаклю описание кинофильма не
    предлагается. Новый текст должен быть длиннее старого (а если старый не обрыв и не «не сюжет» — на 40+ символов)."""
    from xlsx_to_json import norm, title_keys
    index = {}
    with gzip.open(path, "rt", encoding="utf-8") as f:
        for line in f:
            g = json.loads(line)
            if g.get("annotation") and g.get("year") and (not g.get("country") or SOVIET_COUNTRY.search(g["country"])):
                for k in title_keys(g["title"]):
                    index.setdefault(k, []).append(g)
    out, ambiguous = {}, 0
    for film in films:
        y = film["year"]
        cands = {g["gff_id"]: g for k in title_keys(film["title"]) for g in index.get(k, [])
                 if y and abs(g["year"] - y) <= 1}
        exact = [g for g in cands.values() if g["year"] == y]
        pick = exact or list(cands.values())
        if len(pick) != 1:
            ambiguous += len(pick) > 1
            continue
        g = pick[0]
        # каталог ГФФ — кинофильмы: у телеспектакля с тем же названием это обычно другая постановка той же пьесы
        if film["kind"] == "фильм-спектакль" and "спектакл" not in (g.get("genre") or "").lower():
            continue
        new = clean_annotation(g["annotation"])
        if not better(film, new, classes):
            continue
        note = f"ГФФ: «{g['title']}» {g['year']}" + ("" if g["year"] == y else " — год отличается, проверить")
        out[film["film_id"]] = (new, g["url"], note)
    return out, ambiguous


MATCH_PROMPT = """Ниже пары с одинаковым названием и близким годом: аннотация из нашей базы и описание из \
другого источника. Для каждой пары реши, об одном ли фильме тексты.

"yes" — сюжеты согласуются: те же герои, место, время или события. Мелкие расхождения (другое имя у одного \
героя, год, деталь) и разная подробность — это всё равно "yes".
"no" — это РАЗНЫЕ сюжеты: другие герои И другие события; либо один — концерт или документальный фильм, \
а другой — игровой; либо это разные постановки (фильм и телеспектакль).
"unsure" — один из текстов слишком общий или не о сюжете (тема, жанр, судьба фильма), и судить нельзя.

Верни СТРОГО JSON-массив без пояснений вне него: [{{"id": <номер>, "why": "<до 15 слов>", \
"same": "yes"|"no"|"unsure"}}, ...]

{items}"""


def verify_matches(films_by_id, suggest):
    """Сверка предложений моделью FLAG_MODEL: {film_id: "yes"|"no"|"unsure"}; кеш — MATCHES_PATH."""
    cache = json.loads(MATCHES_PATH.read_text(encoding="utf-8")) if MATCHES_PATH.exists() else {}
    key = lambda fid: f"{fid}:{ann_hash(films_by_id[fid]['annotation'])}:{ann_hash(suggest[fid][0])}"
    todo = [fid for fid in suggest if key(fid) not in cache]
    if todo:
        client = get_openai_client()
        for start in range(0, len(todo), 15):
            batch = todo[start:start + 15]
            items = "\n\n".join(
                f"[{fid}] «{films_by_id[fid]['title']}» ({films_by_id[fid]['year']}, {films_by_id[fid]['kind']}"
                f"{', ' + films_by_id[fid]['studio'] if films_by_id[fid]['studio'] else ''})\n"
                f"НАША: {(films_by_id[fid]['annotation'] or '(нет)')[:FLAG_CLIP]}\n"
                f"ДРУГАЯ ({suggest[fid][2]}): {suggest[fid][0][:FLAG_CLIP]}" for fid in batch)
            for attempt in range(3):
                try:
                    resp = client.chat.completions.create(
                        model=FLAG_MODEL, messages=[{"role": "user", "content": MATCH_PROMPT.format(items=items)}],
                        reasoning_effort=FLAG_EFFORT, max_completion_tokens=8000)
                    text = resp.choices[0].message.content.strip().replace("```json", "").replace("```", "")
                    for it in json.loads(text):
                        if int(it["id"]) in batch and it.get("same") in ("yes", "no", "unsure"):
                            cache[key(int(it["id"]))] = it["same"]
                    break
                except Exception as e:      # сбой сети или формата ответа — повтор
                    print(f"  сверка: {e}", file=sys.stderr)
            MATCHES_PATH.write_text(json.dumps(cache, ensure_ascii=False, indent=0), encoding="utf-8")
    return {fid: cache.get(key(fid), "unsure") for fid in suggest}


def cmd_export(args):
    films = load_films(args.db)
    classes = load_classes()
    todo = [f for f in films if f["annotation"] and cached_level(classes, f) is None]
    t_in, t_out, cost = estimate(todo)
    print(f"Фильмов film_id ≥ {FIRST_OPEN_ID}: {len(films)}; без оценки полноты сюжета: {len(todo)}")
    if todo:
        print(f"  классификация {FLAG_MODEL}: ≈ {t_in:,.0f} входных и {t_out:,.0f} выходных токенов, ≈ ${cost:.2f}")
    if args.estimate:
        return
    if args.classify and todo:
        classify(todo, classes)
    suggest, ambiguous = gff_suggestions(films, classes, args.gff) if args.gff else ({}, 0)
    if args.gff:
        print(f"Каталог ГФФ: кандидатов {len(suggest)}; пропущено из-за нескольких кандидатов: {ambiguous}")
    if args.wiki:
        wiki, skipped_series = wiki_suggestions(films, classes, args.wiki, set(suggest))
        print(f"Википедия: кандидатов {len(wiki)}; посерийных сюжетов сериалов со сквозным сюжетом "
              f"(не предлагаются): {skipped_series}")
        suggest.update(wiki)
    if args.web:
        web = web_suggestions(films, classes, args.web, set(suggest), args.web_batches, args.web_skip)
        print(f"Поиск в интернете: кандидатов {len(web)}")
        suggest.update(web)
    # аннотации с латинскими двойниками и «_x000D_» без другой замены — исправленный текст с прежним источником
    fixes = 0
    for f in films:
        if f["film_id"] not in suggest and f["annotation"] and fix_text(f["annotation"]) != f["annotation"]:
            suggest[f["film_id"]] = (clean_annotation(f["annotation"]), f["source"] or "—",
                                     "исправление: латинские буквы в русских словах")
            fixes += 1
    if fixes:
        print(f"Исправлений латинских букв в русских словах: {fixes}")
    if suggest:
        to_verify = {fid: v for fid, v in suggest.items() if not v[2].startswith(("исправление", "поиск в интернете"))}
        verdict = verify_matches({f["film_id"]: f for f in films}, to_verify) if to_verify else {}
        verdict.update({fid: "yes" for fid in suggest if fid not in to_verify})
        dropped = [fid for fid, v in verdict.items() if v == "no"]
        for fid, v in verdict.items():
            text, url, note = suggest[fid]
            suggest[fid] = ((None, None, note + " — по сверке другой фильм, не предложено") if v == "no"
                            else (text, url, note + (" — сверка не уверена, проверить" if v == "unsure" else "")))
        print(f"Сверка моделью: тот же фильм {sum(v == 'yes' for v in verdict.values())}, "
              f"не уверена {sum(v == 'unsure' for v in verdict.values())}, другой фильм {len(dropped)}")
    if suggest:
        for src in ("ГФФ", "Википедия", "поиск в интернете", "исправление"):
            vals = [v for v in suggest.values() if v[0] and v[2].startswith(src)]
            print(f"Предложено: {src} — {len(vals)} (с пометкой «проверить»: "
                  f"{sum(1 for v in vals if 'проверить' in v[2])})")

    rows = []
    for f in films:
        ann = f["annotation"] or ""
        plot = cached_plot(classes, f)
        level = cached_level(classes, f)
        trunc = is_truncated(ann)
        group = ("обрыв" if trunc else "не сюжет" if plot is False else "только завязка" if level == "setup"
                 else "короткая" if len(ann) < SHORT_CHARS else "")
        rank = (0 if trunc else 1 if plot is False else 2 if level == "setup" else 3, len(ann), f["film_id"])
        rows.append((rank, group, f, ann, trunc, level))
    rows.sort(key=lambda r: r[0])

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Аннотации"
    ws.append(HEADERS)
    for n, (_, group, f, ann, trunc, level) in enumerate(rows, start=1):
        query = urllib.parse.quote(f"{f['title']} {f['year'] or ''} фильм".strip())
        ws.append([n, group, f["film_id"], f["title"], f["year"], f["kind"], f["studio"], len(ann),
                   "да" if trunc else "", LEVEL_NAMES.get(level, "?" if ann else ""),
                   ann, f["source"], f'=HYPERLINK("https://ru.wikipedia.org/w/index.php?search={query}","искать")',
                   *suggest.get(f["film_id"], (None, None, None))])
    wrap = Alignment(wrap_text=True, vertical="top")
    top = Alignment(vertical="top")
    for row in ws.iter_rows(min_row=2):
        for cell in row:
            cell.alignment = wrap if cell.column in (4, 11, 14) else top
    for cell in ws[1]:
        cell.font = Font(bold=True)
    for i, w in enumerate(WIDTHS, start=1):
        ws.column_dimensions[openpyxl.utils.get_column_letter(i)].width = w
    ws.freeze_panes = "E2"
    ws.auto_filter.ref = ws.dimensions

    out = Path(args.out)
    out.parent.mkdir(parents=True, exist_ok=True)
    wb.save(out)
    groups = {}
    for r in rows:
        groups[r[1] or "остальные"] = groups.get(r[1] or "остальные", 0) + 1
    unknown = sum(1 for r in rows if r[5] is None)
    print("Группы: " + ", ".join(f"{k} {v}" for k, v in groups.items())
          + (f"; не классифицировано: {unknown}" if unknown else ""))
    print(f"Список работ: {out}")


# Латинские двойники кириллических букв (в каталогах встречается «капpизов» с латинской p — такое слово не
# находится поиском по словам). I/i не трогаем: римские цифры («ХIII»).
HOMOGLYPHS = str.maketrans("aceopxyACEHKMOPTXB", "асеорхуАСЕНКМОРТХВ")


def fix_text(text):
    """Латинские двойники в словах с кириллицей → кириллица; артефакт Excel «_x000D_» → перевод строки."""
    text = str(text or "").replace("_x000D_", "\n")
    return re.sub(r"\w+", lambda m: m.group(0).translate(HOMOGLYPHS)
                  if re.search(r"[а-яё]", m.group(0), re.I) and re.search(r"[a-z]", m.group(0), re.I) else m.group(0),
                  text)


def clean_annotation(text):
    text = fix_text(text).replace("\r\n", "\n").replace("\r", "\n").replace("\xa0", " ")
    text = re.sub(r"[ \t]+", " ", text)
    text = re.sub(r" *\n *", "\n", text)
    return re.sub(r"\n{3,}", "\n\n", text).strip()


def cmd_apply(args):
    db = {f["film_id"]: f for f in load_films(args.db)}
    conn = sqlite3.connect(f"file:{args.db}?mode=ro", uri=True)
    max_id = conn.execute("SELECT MAX(film_id) FROM films").fetchone()[0]
    conn.close()

    edits, errors, warnings, seen = [], [], [], {}
    unchanged = 0
    for path in args.xlsx:
        ws = openpyxl.load_workbook(path, read_only=True).worksheets[0]
        rows = ws.iter_rows(values_only=True)
        header = [str(h or "").strip().lower() for h in next(rows)]

        def col(prefix):
            return next((i for i, h in enumerate(header) if h.startswith(prefix)), None)
        c_id, c_title = col("film_id"), col("название")
        c_new, c_src = col("новая аннотация"), col("источник новой")
        if None in (c_id, c_new, c_src):
            sys.exit(f"{path}: нужны колонки film_id, «Новая аннотация», «Источник новой аннотации»")
        for n, r in enumerate(rows, start=2):
            new = clean_annotation(r[c_new] if c_new < len(r) else None)
            src = str((r[c_src] if c_src < len(r) else None) or "").strip()
            if not new and not src:
                continue
            where = f"{Path(path).name}, строка {n}"
            try:
                fid = int(r[c_id])
            except (TypeError, ValueError):
                errors.append(f"{where}: film_id «{r[c_id]}» — не число")
                continue
            if fid < FIRST_OPEN_ID:
                errors.append(f"{where}: film_id {fid} — каталог ГФФ, аннотации 1–{FIRST_OPEN_ID - 1} не меняются")
                continue
            film = db.get(fid)
            if film is None:
                errors.append(f"{where}: film_id {fid} нет в базе (последний — {max_id})")
                continue
            title = str(r[c_title] or "").strip() if c_title is not None else None
            if title is not None and title != (film["title"] or "").strip():
                errors.append(f"{where}: film_id {fid} в базе — «{film['title']}», в строке — «{title}» "
                              "(строки сдвинуты?)")
                continue
            if fid in seen:
                errors.append(f"{where}: film_id {fid} уже был в {seen[fid]}")
                continue
            seen[fid] = where
            if not new:
                errors.append(f"{where}: film_id {fid} — источник есть, а новой аннотации нет")
                continue
            if not src:
                errors.append(f"{where}: film_id {fid} — не указан источник новой аннотации")
                continue
            if is_truncated(new):
                errors.append(f"{where}: film_id {fid} — новая аннотация обрывается на многоточие")
                continue
            old = film["annotation"] or ""
            if new == old.strip():
                unchanged += 1
                continue
            if len(new) < len(old):
                warnings.append(f"{where}: film_id {fid} «{film['title']}» — новая аннотация короче старой "
                                f"({len(new)} против {len(old)} символов)")
            edits.append({"film_id": fid, "annotation": new, "source": src})

    for w in warnings:
        print("Внимание: " + w)
    for e in errors:
        print("Ошибка: " + e)
    print(f"Правок: {len(edits)}; совпадают с базой: {unchanged}; ошибок: {len(errors)}")
    if errors and not args.skip_errors:
        sys.exit("JSON не записан: исправьте ошибки или запустите с --skip-errors")
    if not edits:
        return
    out = Path(args.out)
    out.parent.mkdir(parents=True, exist_ok=True)
    out.write_text(json.dumps(edits, ensure_ascii=False, indent=1), encoding="utf-8")
    print(f"JSON: {out}\nДальше: python film_search.py import {out} --dry-run")


def main():
    ap = argparse.ArgumentParser(description=__doc__.split("\n")[0])
    ap.add_argument("--db", default=str(ROOT / DB_PATH))
    sub = ap.add_subparsers(dest="cmd", required=True)
    p = sub.add_parser("export", help="выгрузить список работ xlsx")
    p.add_argument("--out", default=str(ROOT / "imports" / "аннотации_список_работ.xlsx"))
    p.add_argument("--classify", action="store_true",
                   help=f"классифицировать «сюжет / не сюжет» моделью {FLAG_MODEL} (платно, см. --estimate)")
    p.add_argument("--estimate", action="store_true", help="только оценить стоимость классификации")
    p.add_argument("--gff", help="каталог gosfilmofond.ru (imports/gff_site_catalog.jsonl.gz) — заполнить предложения")
    p.add_argument("--wiki", help="сюжеты Википедии (imports/wiki_plots.jsonl.gz) — где нет предложения ГФФ")
    p.add_argument("--web", nargs="*", default=[], help="пересказы, найденные поиском в интернете (JSON Lines, .gz или нет)")
    p.add_argument("--web-batches", nargs="*", default=[], help="входные пакеты агентов (JSON) — для проверки хеша аннотации")
    p.add_argument("--web-skip", help="JSON {film_id: причина} — пересказы, исключённые при ручной проверке отчётов агентов")
    p = sub.add_parser("apply", help="заполненный xlsx → JSON-правки для film_search.py import")
    p.add_argument("xlsx", nargs="+")
    p.add_argument("--out", required=True)
    p.add_argument("--skip-errors", action="store_true", help="записать правки без строк с ошибками")
    args = ap.parse_args()
    {"export": cmd_export, "apply": cmd_apply}[args.cmd](args)


if __name__ == "__main__":
    main()
