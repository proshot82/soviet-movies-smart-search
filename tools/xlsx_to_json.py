#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""Сверка списков фильмов в xlsx с базой и подготовка JSON для `film_search.py import`.

    python tools/xlsx_to_json.py список1.xlsx [список2.xlsx …] --out imports/новые.json
    python film_search.py import imports/новые.json --dry-run      # посмотреть, что будет
    python film_search.py import imports/новые.json                # записать
    python film_search.py embed --resume                           # эмбеддинги новых фильмов

Колонки xlsx ищутся по началу заголовка: Название, Жанр, Год, Студия, Аннотация — обязательные
по смыслу; Тип (уточнённый) / Тип, Страна, Режиссёр, Серий, Премьера, ТВ-премьера, Источник, Раздел,
Короткометражный — если есть.

Что делает:
  1. Сводит строки всех файлов: одна и та же запись (название + год + начало аннотации) из разных
     файлов становится одной; берётся самая длинная аннотация, пустые поля дополняются из других.
  2. Сверяет с базой. Фильм считается уже имеющимся, если в базе есть фильм того же года с тем же
     названием (целиком или без уточнения в скобках) и совпадает начало аннотации или студия.
     Одно название и год при разных студиях и аннотациях — это разные фильмы (телеспектакль и
     кинофильм «Старомодная комедия» 1978), такие добавляются.
  3. Тип (kind) — из колонки «Тип (уточнённый)»; если её нет — по жанру и аннотации.
     С --fill-kind тип проставляется и фильмам, которые уже есть в базе без типа (только поле kind).
  4. Пишет JSON (новые фильмы и правки типа по film_id) и отчёт xlsx: решение по каждой строке.
"""
import argparse
import json
import re
import sqlite3
import sys
from collections import defaultdict
from pathlib import Path

import openpyxl

ROOT = Path(__file__).resolve().parent.parent
sys.path.insert(0, str(ROOT))
from film_search import DB_PATH, KINDS  # noqa: E402

# Заголовок (начало, без регистра) → поле. Порядок важен: «Годы» раньше «Год», «Тип (уточн» раньше «Тип».
COLUMNS = [
    ("название", "title"), ("жанр", "genre"), ("годы", None), ("год", "year"), ("студия", "studio"),
    ("аннотация", "annotation"), ("тип (уточн", "kind"), ("тип:", None), ("тип", "kind_raw"),
    ("страна", "country"), ("режисс", "director"), ("серий", "episodes"), ("премьера", "premiere"),
    ("тв-премьера", "tv_premiere"), ("источник", "url"), ("раздел", "section"), ("происхождение", None),
    ("короткометраж", "short"),
]
STUDIO_STOP = {"киностудия", "студия", "студии", "фильм", "фильма", "имени", "творческое", "объединение",
               "главная", "редакция", "программ", "литературно", "драматических", "центральное", "телевидение",
               "телевидения", "кинохроники", "художественных", "фильмов", "союз", "ссср"}


def norm(t):
    t = str(t or "").lower().replace("ё", "е")
    t = re.sub(r"[«»\"„“”'`]", "", t)
    return " ".join(re.sub(r"[^\w]+", " ", t).split())


def title_keys(t):
    """Название целиком и без уточнения в скобках; варианты в кавычках-ёлочках внутри скобок
    («Подполье («В подполье»)») — тоже. Скобки без кавычек («(киноальманах)») вариантом не считаются."""
    keys = {norm(t), norm(re.split(r"\s*[(\[]", str(t or ""), maxsplit=1)[0])}
    for m in re.findall(r"\(([^)]*«[^)]*)\)", str(t or "")):
        keys.update(norm(x) for x in re.findall(r"«([^»]+)»", m))
    return {k for k in keys if k}


def ann_key(a):
    return norm(a)[:60]


def studio_tokens(s):
    return {w for w in norm(s).split() if len(w) > 3 and w not in STUDIO_STOP}


def infer_kind(genre, annotation, studio, default):
    g, a = norm(genre), norm(annotation)[:150]
    if "опер" in g or "балет" in g or (re.search(r"\b(опера|оперы|оперу|балет|балета)\b", a[:60])
                                         and not re.match(r"(фильм )?концерт", a)):
        return "фильм-опера/балет"
    if "мульт" in g or "мульт" in norm(studio) or "мультфильм" in a[:40]:
        return "мультфильм"
    if "концерт" in g or "ревю" in g or re.match(r"(фильм )?(концерт|ревю)", a) or "киноконцерт" in a[:40]:
        return "не игровое"
    if "спектакл" in g or "телеспектакл" in a[:40]:
        return "фильм-спектакль"
    if "телефильм" in a[:40] or "телесериал" in a[:40] or "телевизионн" in a[:40]:
        return "телефильм"
    return default


def read_xlsx(path):
    ws = openpyxl.load_workbook(path, read_only=True).worksheets[0]
    rows = ws.iter_rows(values_only=True)
    header = [str(h or "").strip() for h in next(rows)]
    fields = []
    for h in header:
        hl = h.lower()
        fields.append(next((f for prefix, f in COLUMNS if hl.startswith(prefix)), None))
    if "title" not in fields or "year" not in fields:
        sys.exit(f"{path}: не найдены колонки «Название» и «Год» (заголовки: {header})")
    out = []
    for n, r in enumerate(rows, start=2):
        rec = {f: v for f, v in zip(fields, r) if f and v not in (None, "")}
        if not rec.get("title"):
            continue
        rec = {k: (v.strip() if isinstance(v, str) else v) for k, v in rec.items()}
        try:
            rec["year"] = int(str(rec.get("year"))[:4])
        except ValueError:
            rec["year"] = None
        rec["_file"], rec["_row"] = Path(path).name, n
        out.append(rec)
    return out


def to_film(rec):
    """Строка xlsx → запись для film_search.py import."""
    film = {"title": str(rec["title"]), "year": rec.get("year")}
    for k in ("genre", "studio", "annotation", "country", "director"):
        if rec.get(k) is not None:
            film[k] = str(rec[k])
    if rec.get("episodes") is not None:
        try:
            film["episodes"] = int(str(rec["episodes"]).split()[0])
        except ValueError:
            pass
    dates = [str(rec["premiere"])] if rec.get("premiere") else []
    if rec.get("tv_premiere"):
        dates.append(f"ТВ: {rec['tv_premiere']}")
    if dates:
        film["release_date"] = "; ".join(dates)
    if str(rec.get("short", "")).lower() in ("да", "yes", "1", "true"):
        film["format"] = "короткометражный"
    film["kind"] = rec["kind"]
    film["source"] = rec.get("url") or f"{rec['_file']}" + (f": {rec['section']}" if rec.get("section") else "")
    return film


def main():
    ap = argparse.ArgumentParser(description=__doc__.split("\n")[0])
    ap.add_argument("xlsx", nargs="+")
    ap.add_argument("--db", default=str(ROOT / DB_PATH))
    ap.add_argument("--out", required=True, help="JSON для film_search.py import; отчёт — рядом, .xlsx")
    ap.add_argument("--default-kind", default="телефильм", choices=KINDS,
                    help="тип строк без колонки типа, если по жанру и аннотации не определился")
    ap.add_argument("--fill-kind", action="store_true",
                    help="проставить тип фильмам, которые уже есть в базе без типа (только поле kind)")
    args = ap.parse_args()

    # ── 1. строки всех файлов, сведённые в записи ──
    rows = [r for path in args.xlsx for r in read_xlsx(path)]
    for r in rows:
        k = (r.get("kind") or "").strip()
        r["kind"] = k if k in KINDS else infer_kind(r.get("genre"), r.get("annotation"), r.get("studio"),
                                                    args.default_kind)
    records, by_key = [], {}
    for r in rows:
        key = (norm(r["title"]), r["year"], ann_key(r.get("annotation")))
        if key in by_key:
            rec = by_key[key]
            rec["_rows"].append(r)
            if len(str(r.get("annotation") or "")) > len(str(rec.get("annotation") or "")):
                rec["annotation"] = r["annotation"]
            for f, v in r.items():
                if not f.startswith("_") and rec.get(f) in (None, ""):
                    rec[f] = v
        else:
            rec = dict(r, _rows=[r])
            by_key[key] = rec
            records.append(rec)

    # ── 2. сверка с базой ──
    conn = sqlite3.connect(f"file:{args.db}?mode=ro", uri=True)
    cols = {r[1] for r in conn.execute("PRAGMA table_info(films)")}
    kind_sql = "kind" if "kind" in cols else "NULL"
    db = conn.execute(f"SELECT film_id, title, year, studio, annotation, {kind_sql} FROM films").fetchall()
    db_index = defaultdict(list)
    for row in db:
        for k in title_keys(row[1]):
            db_index[(k, row[2])].append(row)

    def find_in_db(rec):
        cands = {row[0]: row for k in title_keys(rec["title"]) for row in db_index.get((k, rec["year"]), [])}
        for row in cands.values():
            same_ann = ann_key(row[4]) and ann_key(row[4]) == ann_key(rec.get("annotation"))
            if same_ann or studio_tokens(row[3]) & studio_tokens(rec.get("studio")):
                return row, list(cands)
        return None, list(cands)

    films, report, kind_updates = [], [], {}
    for rec in records:
        found, cands = find_in_db(rec)
        if found:
            decision, film_id = "уже в базе", found[0]
            if args.fill_kind and found[5] is None and rec["kind"]:
                kind_updates.setdefault(film_id, rec["kind"])
        else:
            decision, film_id = "добавить", None
            films.append(to_film(rec))
        note = ("в базе есть фильм с тем же названием и годом (film_id "
                + ", ".join(map(str, cands)) + ") — другой: не совпали ни студия, ни аннотация") \
            if cands and not found else ""
        for i, r in enumerate(rec["_rows"]):
            report.append([r["_file"], r["_row"], r["title"], r["year"], rec["kind"],
                           decision if i == 0 else "та же запись, что выше (другой файл)",
                           film_id, len(str(rec.get("annotation") or "")), note if i == 0 else ""])

    if args.fill_kind:
        # фильмы базы без типа, не встретившиеся в списках: тип по жанру и аннотации, по умолчанию — кинофильм
        for fid, title, year, studio, ann, kind in db:
            if kind is None and fid not in kind_updates:
                genre = conn.execute("SELECT genre FROM films WHERE film_id = ?", (fid,)).fetchone()[0]
                kind_updates[fid] = infer_kind(genre, ann, studio, "кинофильм")

    out = Path(args.out)
    out.parent.mkdir(parents=True, exist_ok=True)
    payload = films + [{"film_id": fid, "kind": k} for fid, k in sorted(kind_updates.items())]
    out.write_text(json.dumps(payload, ensure_ascii=False, indent=1), encoding="utf-8")

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Сверка"
    ws.append(["Файл", "Строка", "Название", "Год", "Тип", "Решение", "film_id в базе", "Длина аннотации",
               "Примечание"])
    for r in report:
        ws.append(r)
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = ws.dimensions
    for col, w in zip("ABCDEFGHI", (28, 8, 45, 7, 16, 32, 12, 10, 70)):
        ws.column_dimensions[col].width = w
    rep_path = out.with_name(out.stem + "_отчёт.xlsx")
    wb.save(rep_path)

    n_rows = len(rows)
    in_db = sum(1 for r in report if r[5] == "уже в базе")
    print(f"Строк во входных файлах: {n_rows}; разных записей: {len(records)}")
    print(f"  уже в базе: {in_db}; к добавлению: {len(films)}; правок типа у фильмов базы: {len(kind_updates)}")
    kinds = defaultdict(int)
    for f in films:
        kinds[f["kind"]] += 1
    print("  добавляются по типам: " + ", ".join(f"{k} {v}" for k, v in sorted(kinds.items(), key=lambda x: -x[1])))
    print(f"JSON: {out}\nОтчёт: {rep_path}")


if __name__ == "__main__":
    main()
